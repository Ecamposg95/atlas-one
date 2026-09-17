"""Exportar/importar catalogo (Excel/CSV) con columnas Color y Talla.

Export: una fila por variante, con Color/Talla al final (orden existente
intacto). Upload: filas con el mismo Nombre*/Color/Talla en la misma carga
se agrupan en un producto con N variantes via `crear_variantes`.
"""
import io

from decimal import Decimal

from conftest import _make_product

from app.models.inventory import StockOnHand
from app.models.products import Product, ProductVariant


def _h(auth, org):
    return {**auth, "X-Organization-ID": str(org.id)}


class TestExportConVariantes:
    def test_una_fila_por_variante_con_color_y_talla(self, client, db, org, branch_a, auth_admin):
        p, principal = _make_product(db, org, "Playera lisa", "PLY-R-S", 120, [(branch_a.id, True)])
        principal.color, principal.size, principal.variant_name = "Rojo", "S", "Rojo / S"
        hermana = ProductVariant(
            product_id=p.id, sku="PLY-R-M", price=Decimal("120"), cost=Decimal("70"),
            color="Rojo", size="M", variant_name="Rojo / M", organization_id=org.id,
        )
        db.add(hermana)
        db.commit()

        r = client.get("/api/products/export/excel", headers=_h(auth_admin, org))
        assert r.status_code == 200, r.text

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert headers[:2] == ["SKU", "Nombre"], "el orden de columnas existente no cambia"
        assert headers[-2:] == ["Color", "Talla"], "Color/Talla se agregan al final"

        filas = {row[headers.index("SKU")]: row for row in ws.iter_rows(min_row=2, values_only=True)}
        assert set(filas) == {"PLY-R-S", "PLY-R-M"}
        assert filas["PLY-R-S"][headers.index("Color")] == "Rojo"
        assert filas["PLY-R-S"][headers.index("Talla")] == "S"
        assert filas["PLY-R-M"][headers.index("Talla")] == "M"


class TestUploadConVariantes:
    def _csv(self, filas, cabeceras):
        buf = io.StringIO()
        buf.write(",".join(cabeceras) + "\n")
        for f in filas:
            buf.write(",".join(str(f.get(c, "")) for c in cabeceras) + "\n")
        return buf.getvalue().encode("utf-8")

    CABECERAS = ["SKU", "Nombre", "Departamento", "Precio Base", "Costo", "Stock", "Color", "Talla"]

    def test_agrupa_filas_del_mismo_nombre_en_un_producto(self, client, db, org, auth_admin):
        filas = [
            {"SKU": "PLY-R-S", "Nombre": "Playera lisa", "Departamento": "Playeras",
             "Precio Base": "120", "Costo": "60", "Stock": "3", "Color": "Rojo", "Talla": "S"},
            {"SKU": "PLY-R-M", "Nombre": "Playera lisa", "Departamento": "Playeras",
             "Precio Base": "120", "Costo": "60", "Stock": "5", "Color": "Rojo", "Talla": "M"},
            {"SKU": "PLY-N-M", "Nombre": "Playera lisa", "Departamento": "Playeras",
             "Precio Base": "130", "Costo": "60", "Stock": "0", "Color": "Negro", "Talla": "M"},
        ]
        content = self._csv(filas, self.CABECERAS)
        r = client.post(
            "/api/products/upload",
            headers=_h(auth_admin, org),
            files={"file": ("ropa.csv", content, "text/csv")},
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["created"] == 3

        productos = db.query(Product).filter(Product.organization_id == org.id, Product.name == "Playera lisa").all()
        assert len(productos) == 1, "las 3 filas deben agruparse en un solo producto"
        p = productos[0]
        vs = {v.sku: v for v in db.query(ProductVariant).filter(ProductVariant.product_id == p.id).all()}
        assert set(vs) == {"PLY-R-S", "PLY-R-M", "PLY-N-M"}
        assert vs["PLY-R-M"].variant_name == "Rojo / M"
        assert float(vs["PLY-N-M"].price) == 130.0

        stock = {sku: db.query(StockOnHand).filter(StockOnHand.variant_id == v.id).first() for sku, v in vs.items()}
        assert float(stock["PLY-R-M"].qty_on_hand) == 5.0

    def test_pareja_repetida_entre_hermanas_no_se_pierde_por_cache(self, client, db, org, auth_admin):
        """3+ filas del mismo grupo: la deteccion de pareja repetida debe ver
        TODAS las hermanas ya creadas en filas previas de esta carga, no solo
        la principal (`producto.variants` se cachea en el primer acceso)."""
        filas = [
            {"SKU": "GOR-R-S", "Nombre": "Gorra lisa", "Departamento": "Gorras",
             "Precio Base": "80", "Costo": "40", "Stock": "1", "Color": "Rojo", "Talla": "S"},
            {"SKU": "GOR-R-M", "Nombre": "Gorra lisa", "Departamento": "Gorras",
             "Precio Base": "80", "Costo": "40", "Stock": "1", "Color": "Rojo", "Talla": "M"},
            {"SKU": "GOR-R-M-2", "Nombre": "Gorra lisa", "Departamento": "Gorras",
             "Precio Base": "80", "Costo": "40", "Stock": "1", "Color": "rojo", "Talla": "m"},
        ]
        content = self._csv(filas, self.CABECERAS)
        r = client.post(
            "/api/products/upload",
            headers=_h(auth_admin, org),
            files={"file": ("gorras.csv", content, "text/csv")},
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["created"] == 2
        assert body["failed"] == 1
        skus = {v.sku for v in db.query(ProductVariant).join(Product).filter(
            Product.organization_id == org.id, Product.name == "Gorra lisa").all()}
        assert skus == {"GOR-R-S", "GOR-R-M"}
