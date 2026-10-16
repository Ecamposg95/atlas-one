"""``GET /api/products/export/excel`` and ``POST /api/products/upload``
— bulk Excel/CSV catalog import-export.

Sprint 5b split — extracted verbatim from the original ``products.py``.
"""
from __future__ import annotations

import csv
import io
import logging
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session, joinedload

from app.core.database import get_db
from app.core.permissions import require_module
from app.core.security import get_current_user
from app.core.tenant_context import get_current_active_organization
from app.crud.products import _is_admin, query_visible_products
from app.models import Brand, Department, Product, ProductVariant, StockOnHand, User
from app.models.organization import Branch, BranchType
from app.modules.products.services.importacion import ContextoImportacion, FilaInvalida, importar_fila
from app.modules.products.services.validaciones import validar_sucursales_de_la_org
from app.services.barcodes import AsignadorDeCodigos

from ._shared import _PRODUCT_ADVANCED_ROLES

logger = logging.getLogger(__name__)

router = APIRouter()


# -----------------------------
# 6. Exportar Excel (Template con productos existentes)
# -----------------------------
@router.get("/export/excel")
def export_products_template(
    branch_id: Optional[int] = None,
    search: Optional[str] = None,
    department_id: Optional[str] = None,
    brand_id: Optional[str] = None,
    approval_status: Optional[str] = None,
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    org_id: int = Depends(get_current_active_organization)
):
    """
    Generates an Excel template with existing products for bulk update/import.

    Acepta los mismos filtros que `GET /products/` para que el "Exportar"
    del Admin Catalog descargue exactamente el subconjunto que el admin
    está viendo en pantalla (no el catálogo completo).

    [A1-08] Usa query_visible_products como fuente — CAJERO/GERENTE no
    pueden exportar catálogo cross-branch vía este endpoint.
    """
    import io
    from fastapi.responses import StreamingResponse

    # Track 3: cajero puede exportar SU sucursal con include_inactive.
    # branch_override (drill-down a otra sucursal) sigue admin-only —
    # query_visible_products limita a la sucursal del usuario para no-admin.
    branch_override = None
    if _is_admin(current_user) and branch_id and branch_id != 0:
        branch_override = branch_id

    q = query_visible_products(
        db,
        current_user,
        org_id,
        include_inactive=include_inactive,
        search=search or None,
        branch_id_override=branch_override,
    ).options(
        joinedload(Product.variants).joinedload(ProductVariant.prices),
        joinedload(Product.variants).joinedload(ProductVariant.packaging_units),
        joinedload(Product.department),
    ).filter(Product.deleted_at == None)

    if department_id:
        q = q.filter(Product.department_id == department_id)
    if brand_id:
        q = q.filter(Product.brand_id == brand_id)
    if approval_status and approval_status != 'ALL':
        q = q.filter(Product.approval_status == approval_status)

    query = q.all()

    data = []

    # 1. Resolve Target Branch (Explicit OR User's Branch OR Org's HQ)
    target_branch_id = branch_id
    if not target_branch_id:
        target_branch_id = current_user.branch_id

    if not target_branch_id:
        hq = db.query(Branch).filter(
            Branch.organization_id == org_id,
            Branch.branch_type == BranchType.HQ
        ).first()
        if hq:
            target_branch_id = hq.id
        else:
            # Fallback to any branch if HQ not found
            any_b = db.query(Branch).filter(Branch.organization_id == org_id).first()
            if any_b: target_branch_id = any_b.id

    # 2. Add existing products — una fila por variante (color/talla incluidos)
    for p in query:
        if not p.variants: continue
        for v in p.variants:
            if v.deleted_at is not None:
                # Una talla retirada no debe exportarse: si el archivo se
                # vuelve a subir, la busqueda por SKU/el indice unico ignoran
                # las filas soft-deleted y la re-crearian como duplicado.
                continue

            # Get Stock
            stock_qty = 0
            if target_branch_id:
                 s = db.query(StockOnHand).filter(
                     StockOnHand.variant_id == v.id,
                     StockOnHand.branch_id == target_branch_id
                 ).first()
                 if s: stock_qty = s.qty_on_hand

            row = {
                "SKU": v.sku,
                "Nombre": p.name,
                "Descripcion": p.description or "",
                "Unidad": p.unit,
                "Departamento": p.department.name if p.department else "General",
                "Marca": p.brand.name if p.brand else "",
                "Codigo Barras": v.barcode or "",
                "Precio Base": float(v.price),
                "Costo": float(v.cost),
                "Stock": float(stock_qty),
                "Incluye IVA": "Si" if v.has_iva else "No",
                "Color": v.color or "",
                "Talla": v.size or "",
            }

            # Prices (up to 5)
            prices_sorted = sorted(v.prices, key=lambda x: x.min_quantity)
            for i, price in enumerate(prices_sorted[:5]):
                idx = i + 1
                row[f"P{idx} Nombre"] = price.price_name
                row[f"P{idx} Min"] = float(price.min_quantity)
                row[f"P{idx} Precio"] = float(price.unit_price)
                # Find linked pkg name if any
                pkg_name = ""
                if price.linked_package_id:
                    # Naive search in this variant's packages
                     pkg = next((pk for pk in v.packaging_units if pk.id == price.linked_package_id), None)
                     if pkg: pkg_name = pkg.name
                row[f"P{idx} Empaque"] = pkg_name

            # Packaging (up to 3)
            pkgs_sorted = sorted(v.packaging_units, key=lambda x: x.units_per_package)
            for i, pkg in enumerate(pkgs_sorted[:3]):
                idx = i + 1
                row[f"E{idx} Nombre"] = pkg.name
                row[f"E{idx} Barcode"] = pkg.barcode or ""
                row[f"E{idx} Cantidad"] = float(pkg.units_per_package)
                row[f"E{idx} Precio"] = float(pkg.package_price)

            data.append(row)

    # 2. Build final column order (no pandas needed)
    standard_cols = [
        "SKU", "Nombre", "Descripcion", "Unidad", "Departamento", "Marca", "Codigo Barras",
        "Precio Base", "Costo", "Stock", "Incluye IVA"
    ]
    final_cols = standard_cols + \
                 [f"P{i} {f}" for i in range(1, 6) for f in ["Nombre", "Min", "Precio", "Empaque"]] + \
                 [f"E{i} {f}" for i in range(1, 4) for f in ["Nombre", "Barcode", "Cantidad", "Precio"]] + \
                 ["Color", "Talla"]

    # 3. Write Excel with openpyxl directly
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    output = io.BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Plantilla'

    # Write header row
    worksheet.append(final_cols)

    # Write data rows (fill missing columns with empty string)
    for row_dict in data:
        worksheet.append([row_dict.get(col, "") for col in final_cols])

    # --- 1. DATA VALIDATION ---
    depts = [d.name for d in db.query(Department).filter(Department.organization_id == org_id).all()]
    brands = [b.name for b in db.query(Brand).filter(Brand.organization_id == org_id).all()]
    units = ["pza", "kg", "g", "ml", "L", "m", "cm", "caja", "pqte", "srv"]

    if depts or brands:
        val_sheet = workbook.create_sheet("Listas_Validacion")
        val_sheet.sheet_state = 'hidden'

        for i, d in enumerate(depts):
            val_sheet.cell(row=i+1, column=1, value=d)
        for i, b in enumerate(brands):
            val_sheet.cell(row=i+1, column=2, value=b)
        for i, u in enumerate(units):
            val_sheet.cell(row=i+1, column=3, value=u)

        max_row = max(1000, len(data) + 100)

        if depts:
            dv_dept = DataValidation(type="list", formula1=f"'Listas_Validacion'!$A$1:$A${len(depts)}", allow_blank=True)
            worksheet.add_data_validation(dv_dept)
            dv_dept.add(f"E2:E{max_row}")

        if brands:
            dv_brand = DataValidation(type="list", formula1=f"'Listas_Validacion'!$B$1:$B${len(brands)}", allow_blank=True)
            worksheet.add_data_validation(dv_brand)
            dv_brand.add(f"F2:F{max_row}")

        dv_unit = DataValidation(type="list", formula1=f"'Listas_Validacion'!$C$1:$C${len(units)}", allow_blank=True)
        worksheet.add_data_validation(dv_unit)
        dv_unit.add(f"D2:D{max_row}")

    # --- 2. STYLING ---
    header_fill = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for col in worksheet.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        worksheet.column_dimensions[column].width = max(10, min(max_length + 2, 50))

    workbook.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    headers = {
        'Content-Disposition': f'attachment; filename="catalogo_productos_{timestamp}.xlsx"'
    }
    return StreamingResponse(output, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers=headers)

async def _leer_filas(file: UploadFile) -> list[dict]:
    """CSV o Excel -> lista de dicts con cabeceras en minúsculas."""
    nombre = (file.filename or "").lower()
    es_csv = nombre.endswith(".csv")
    es_excel = nombre.endswith((".xlsx", ".xlsm", ".xls"))
    if not (es_csv or es_excel):
        raise HTTPException(status_code=400, detail="Formato inválido. Use Excel o CSV.")
    contenido = await file.read()
    if not contenido:
        raise HTTPException(status_code=400, detail="Archivo vacío.")
    try:
        if es_csv:
            texto = contenido.decode("utf-8-sig", errors="replace")
            return [{k.lower().strip(): v for k, v in r.items()} for r in csv.DictReader(io.StringIO(texto))]
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(contenido), data_only=True).active
        cabeceras = [str(c.value or "").lower().strip() for c in ws[1]]
        return [{cabeceras[i]: v for i, v in enumerate(fila)} for fila in ws.iter_rows(min_row=2, values_only=True)]
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="El archivo no se pudo leer. Verifique que sea un Excel o CSV válido e intente de nuevo.",
        )


def _sucursal_hq_o_cualquiera(db: Session, org_id: int) -> Optional[int]:
    hq = db.query(Branch.id).filter(Branch.organization_id == org_id, Branch.branch_type == BranchType.HQ).first()
    if hq:
        return hq.id
    cualquiera = db.query(Branch.id).filter(Branch.organization_id == org_id).first()
    return cualquiera.id if cualquiera else None


def _sucursales_de_carga(db: Session, org_id: int, user: User, scope: str,
                         target_branch_ids: Optional[str]) -> list[int]:
    """Dónde queda habilitado lo importado. Usuario con sucursal: siempre la
    suya (no puede afectar otras). HQ: `scope` all | selected | branch (=HQ)."""
    if user.branch_id:
        return [user.branch_id]
    if scope == "all":
        ids = [b.id for b in db.query(Branch.id).filter(Branch.organization_id == org_id).all()]
    elif scope == "selected" and target_branch_ids:
        try:
            ids = [int(x.strip()) for x in target_branch_ids.split(",") if x.strip()]
        except ValueError:
            raise HTTPException(status_code=400, detail=f"target_branch_ids inválido: '{target_branch_ids}'")
        validar_sucursales_de_la_org(db, org_id, ids)
    else:
        hq = db.query(Branch.id).filter(Branch.organization_id == org_id, Branch.branch_type == BranchType.HQ).first()
        ids = [hq.id] if hq else []
    return sorted(set(ids))


def _contexto_de_carga(db: Session, org_id: int, user: User, branch_id: Optional[int],
                       target_branch_ids: Optional[str], scope: str, filas: list[dict]) -> ContextoImportacion:
    # RBAC: un usuario con sucursal (CAJERO/GERENTE) importa a la suya, se
    # ignora cualquier branch_id del form.
    if user.branch_id:
        destino = user.branch_id
    else:
        destino = branch_id
        if destino and db.query(Branch.id).filter(Branch.id == destino, Branch.organization_id == org_id).first() is None:
            raise HTTPException(status_code=400, detail="Sucursal de destino inválida o no pertenece a tu organización.")
    if not destino:
        destino = _sucursal_hq_o_cualquiera(db, org_id)
    columnas = set(filas[0].keys()) if filas else set()
    return ContextoImportacion(
        org_id=org_id,
        user=user,
        target_branch_id=destino,
        sucursales=_sucursales_de_carga(db, org_id, user, scope, target_branch_ids),
        dept_map={(d.name or "").lower(): d.id for d in db.query(Department).filter(Department.organization_id == org_id).all() if d.name},
        brand_map={(b.name or "").lower(): b.id for b in db.query(Brand).filter(Brand.organization_id == org_id).all() if b.name},
        has_pkg_cols=any(k.startswith("e1") for k in columnas),
        has_tier_cols=any(k.startswith("p1") for k in columnas),
        # Un solo asignador para todo el archivo: no releer el MAX por fila.
        asignador=AsignadorDeCodigos(db, org_id),
    )


def _confirmar(db: Session, hasta_fila: int, confirmada_hasta: int = 0) -> None:
    """Commit por lote. Si falla, se dice: antes se tragaba y el cliente
    recibía "created: 50" con la transacción caída. `confirmada_hasta` es la
    última fila del archivo que un lote anterior ya dejó guardada."""
    try:
        db.commit()
    except Exception:
        db.rollback()
        logger.exception("upload: no se pudo guardar el lote (hasta la fila %s)", hasta_fila)
        previo = (f"Las filas hasta la {confirmada_hasta} ya se habían guardado en lotes anteriores; "
                  if confirmada_hasta else "Ninguna fila se había guardado antes; ")
        raise HTTPException(
            status_code=500,
            detail=f"No se pudieron guardar las filas procesadas (hasta la fila {hasta_fila}). "
                   f"{previo}el resto no se guardó. Revise el archivo e inténtelo de nuevo.",
        )


def _exigir_puede_importar(user: User) -> None:
    if user.role not in _PRODUCT_ADVANCED_ROLES:
        raise HTTPException(status_code=403, detail="No autorizado para importar productos")
    if not _is_admin(user) and not user.branch_id:
        # Mismo texto que `crear_producto`: sin esto todas las filas fallarían una por una.
        raise HTTPException(status_code=403, detail="Solo usuarios con sucursal asignada pueden crear productos.")


def _procesar_filas(db: Session, org_id: int, ctx: ContextoImportacion, filas: list[dict], dry_run: bool) -> dict:
    creadas = actualizadas = fallidas = 0
    detalles: list[dict] = []
    preview: list[dict] = []
    confirmada_hasta = 0
    for idx, fila in enumerate(filas):
        try:
            with db.begin_nested():
                res = importar_fila(db, ctx, fila)
        except Exception as e:   # noqa: BLE001 — cada fila se aísla y se reporta
            detalle = e.detail if isinstance(e, HTTPException) else str(e)
            if isinstance(e, (FilaInvalida, HTTPException)):
                logger.warning("upload org=%s fila %s: %s", org_id, idx + 2, detalle)
            else:
                logger.exception("upload org=%s fila %s: error inesperado", org_id, idx + 2)
            fallidas += 1
            detalles.append({"row": idx + 2, "sku": str(fila.get("sku", "N/A")), "error": detalle})
            if len(preview) < 20:
                preview.append({"action": "ERROR", "sku": fila.get("sku") or fila.get("código") or fila.get("codigo") or None,
                                "name": fila.get("nombre") or fila.get("name"), "price": None, "error_message": detalle})
        else:
            if res.accion == "SKIP":
                continue
            creadas += res.accion == "NEW"
            actualizadas += res.accion == "UPDATE"
            if len(preview) < 20:
                preview.append({"action": res.accion, "sku": res.sku, "name": res.nombre,
                                "price": res.precio, "error_message": None})
        # Commit por lotes para no sostener el lock; en dry_run no se confirma nada.
        if not dry_run and (creadas + actualizadas + fallidas) % 50 == 0:
            _confirmar(db, hasta_fila=idx + 2, confirmada_hasta=confirmada_hasta)
            confirmada_hasta = idx + 2
    return {"creadas": creadas, "actualizadas": actualizadas, "fallidas": fallidas,
            "detalles": detalles, "preview": preview, "confirmada_hasta": confirmada_hasta}


@router.post("/upload", dependencies=[Depends(require_module("catalog"))])
async def upload_products(
    branch_id: Optional[int] = Form(None),
    target_branch_ids: Optional[str] = Form(None),   # lista separada por comas
    scope: str = Form("branch"),                     # "branch" | "all" | "selected"
    dry_run: bool = Form(False),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    org_id: int = Depends(get_current_active_organization),
):
    """Carga masiva. Parseo y sucursales en helpers de este módulo; cada fila
    en `services.importacion.importar_fila` dentro de su SAVEPOINT."""
    _exigir_puede_importar(current_user)
    filas = await _leer_filas(file)
    ctx = _contexto_de_carga(db, org_id, current_user, branch_id, target_branch_ids, scope, filas)
    r = _procesar_filas(db, org_id, ctx, filas, dry_run)
    if dry_run:
        db.rollback()
        return {"total_rows": r["creadas"] + r["actualizadas"] + r["fallidas"], "to_create": r["creadas"],
                "to_update": r["actualizadas"], "errors": r["fallidas"], "preview": r["preview"],
                "error_details": r["detalles"]}
    _confirmar(db, hasta_fila=len(filas) + 1, confirmada_hasta=r["confirmada_hasta"])
    return {"created": r["creadas"], "updated": r["actualizadas"], "failed": r["fallidas"],
            "pending": 0, "details": r["detalles"]}
